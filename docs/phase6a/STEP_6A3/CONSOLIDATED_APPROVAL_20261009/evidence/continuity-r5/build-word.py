from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path('/Users/andresbarouh/repos/tcdx-grc/docs/phase6a/STEP_6A3/CONSOLIDATED_APPROVAL_20261009')
content=json.loads((B/'evidence/continuity-r5/word-content.json').read_text())
d=Document();section=d.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.7);section.bottom_margin=Inches(.65);section.left_margin=Inches(.75);section.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 for border in list(s.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(23)
d.styles['Heading 1'].font.size=Pt(17)
d.styles['Heading 1'].paragraph_format.space_after=Pt(13)
for i,s in enumerate(content):
 if i:d.add_page_break()
 d.add_paragraph(s['title'],'Title' if i==0 else 'Heading 1')
 for text in s['body']:
  p=d.add_paragraph(text)
  p.paragraph_format.widow_control=True
  if text.startswith('[ ]'):
   p.paragraph_format.left_indent=Inches(.1)
   for run in p.runs:run.bold=True
  if i==len(content)-1 and text.startswith('S'):
   p.paragraph_format.keep_together=True
   for run in p.runs:run.font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('R5  DRAFT_NOT_APPROVED  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.name='Arial';run.font.size=Pt(8)
d.core_properties.title='Decisiones macro para continuar STEP 6A3'
d.core_properties.subject='R5 fuentes bindings IANA RLS y preparación QA'
d.core_properties.author='';d.core_properties.last_modified_by=''
out=B/'R5_MACRO_DECISIONS.docx';d.save(out)
print('DOCX_CREATED',out)
